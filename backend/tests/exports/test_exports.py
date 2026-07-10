"""POST /api/runs/{id}/exports (async job -> ExportArtifact) + GET
/api/exports/{id} (presigned URL). Full notebook parity: workbook, both
chart PNGs, and the frequency/coded CSVs, all downloadable via the API."""

import io

import httpx
import openpyxl

TASTE_TEXTS = ["The taste is delicious and sweet.", "I love the flavor.", "Tastes amazing.", "Very tasty and delicious."]
PRICE_TEXTS = ["Great value for the price.", "It's affordable and cheap.", "Good price point.", "Worth the cost."]

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _export(client, run_id, fmt):
    resp = client.post(f"/api/runs/{run_id}/exports", params={"format": fmt})
    assert resp.status_code == 202, resp.text
    job_id = resp.json()["job_id"]
    job = client.get(f"/api/jobs/{job_id}").json()
    assert job["status"] == "finished", job
    artifact_id = job["result"]["artifact_id"]
    return client.get(f"/api/exports/{artifact_id}").json()


def test_xlsx_export_has_five_sheets_and_downloads_via_presigned_url(build_applied_run):
    seeded = build_applied_run(TASTE_TEXTS + PRICE_TEXTS, ["Taste & Flavor", "Price & Value"])
    export = _export(seeded["client"], seeded["run_id"], "xlsx")
    assert export["kind"] == "xlsx"

    r = httpx.get(export["url"])
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) == 5
    assert wb.sheetnames == ["Summary", "Frequencies", "Co-occurrence", "Codebook", "Coded data"]
    assert wb["Coded data"].max_row - 1 == len(TASTE_TEXTS) + len(PRICE_TEXTS)  # -1 for header


def test_csv_frequencies_export_downloads_and_has_expected_rows(build_applied_run):
    seeded = build_applied_run(TASTE_TEXTS + PRICE_TEXTS, ["Taste & Flavor", "Price & Value"])
    export = _export(seeded["client"], seeded["run_id"], "csv_frequencies")

    r = httpx.get(export["url"])
    assert r.status_code == 200
    lines = r.text.strip().splitlines()
    assert lines[0].startswith("code_id,code,count,percentage") or lines[0].startswith("code,count,percentage")
    assert len(lines) - 1 == 3  # Taste & Flavor, Price & Value, Other


def test_csv_coded_export_downloads_and_has_one_row_per_response(build_applied_run):
    seeded = build_applied_run(TASTE_TEXTS + PRICE_TEXTS, ["Taste & Flavor", "Price & Value"])
    export = _export(seeded["client"], seeded["run_id"], "csv_coded")

    r = httpx.get(export["url"])
    assert r.status_code == 200
    lines = r.text.strip().splitlines()
    assert len(lines) - 1 == len(TASTE_TEXTS) + len(PRICE_TEXTS)


def test_png_frequency_export_is_a_valid_image(build_applied_run):
    seeded = build_applied_run(TASTE_TEXTS + PRICE_TEXTS, ["Taste & Flavor", "Price & Value"])
    export = _export(seeded["client"], seeded["run_id"], "png_frequencies")

    r = httpx.get(export["url"])
    assert r.status_code == 200
    assert r.content[:8] == _PNG_MAGIC
    assert len(r.content) > 1024


def test_png_cooccurrence_export_is_a_valid_image(build_applied_run):
    seeded = build_applied_run(TASTE_TEXTS + PRICE_TEXTS, ["Taste & Flavor", "Price & Value"])
    export = _export(seeded["client"], seeded["run_id"], "png_cooccurrence")

    r = httpx.get(export["url"])
    assert r.status_code == 200
    assert r.content[:8] == _PNG_MAGIC
    assert len(r.content) > 1024


def test_export_requires_applied_run(seed_dataset):
    seeded = seed_dataset(TASTE_TEXTS)
    client = seeded["client"]
    draft_run_id = client.post(f"/api/questions/{seeded['question_id']}/runs", json={
        "kind": "import", "codebook_source": ["Taste & Flavor"],
    }).json()["id"]

    resp = client.post(f"/api/runs/{draft_run_id}/exports", params={"format": "xlsx"})
    assert resp.status_code == 409


def test_get_export_unknown_artifact_returns_404(build_applied_run):
    seeded = build_applied_run(TASTE_TEXTS, ["Taste & Flavor"])
    resp = seeded["client"].get("/api/exports/999999")
    assert resp.status_code == 404
